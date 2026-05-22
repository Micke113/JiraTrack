import json
import os
import re
import requests
import time
import xlsxwriter
import yaml
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from requests.auth import HTTPBasicAuth


# --- CONFIGURATION ---
DEFAULT_FIELDS = "customfield_10005, customfield_18002, updated, status, summary, issuetype, components, priority, created"
WEEKS_RANGE = time.strftime("%W", time.gmtime())  # Nombre de semaines depuis le début de l'année

OUTPUT_EXCEL = f"output/jira_tickets_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
EXCEL_TEMPLATE = "xlsx_templates/TimeTracking_PIX_LMI_template.xlsx"

# JQL Request
TICKET_HEADERS = ["Sprint", "Date", "DayLog", "Ticket", "Hours spent", "Title", "Project", "Component", "Type", "Planned/Unplanned", "Cause"]


# --- Gestion credentials ---
def load_credentials() -> dict:
    if os.path.exists("config/credentials.yaml"):
        with open("config/credentials.yaml", "r", encoding="utf-8") as f:
            credential_data = yaml.load(f, Loader=yaml.CSafeLoader)
        return credential_data
    return {}


def save_credentials(data: dict) -> None:
    credentials = load_credentials()
    if not credentials:
        os.makedirs("config", exist_ok=True)

    cookie = data.get('jira_cookie', "").strip()
    cookie = cookie.encode('ascii', 'ignore').decode('ascii')
    cookie = re.sub(r'[^\x20-\x7E]', '', cookie).strip()
    credentials['jira_cookie'] = cookie
    credentials['jira_link'] = data.get('jira_link', "")
    credentials['jira_user'] = data.get('jira_user', "")

    with open("config/credentials.yaml", "w", encoding="utf-8") as f:
        yaml.dump(credentials, f, Dumper=yaml.CSafeDumper, default_flow_style=False, allow_unicode=True)


# --- Gestion du cookie ---
def load_cookie() -> str:
    """Charge le cookie depuis un fichier texte local."""
    credential_data = load_credentials()
    if credential_data and credential_data.get('jira_cookie'):
        cookie = credential_data['jira_cookie'].strip()
        cookie = cookie.encode('ascii', 'ignore').decode('ascii')
        cookie = re.sub(r'[^\x20-\x7E]', '', cookie).strip()
        if cookie:
            print("🍪 Cookie Jira chargé depuis le fichier local.")
            return cookie
    return ""


# --- Sauvegarde du cookie ---
def save_cookie(cookie):
    """Sauvegarde le cookie dans un fichier local."""
    credential_data = load_credentials()
    credential_data["jira_cookie"] = cookie
    with open("config/credentials.yaml", "w", encoding="utf-8") as f:
        yaml.dump(credential_data, f, Dumper=yaml.CSafeDumper, default_flow_style=False, allow_unicode=True)
    print("💾 Nouveau cookie sauvegardé dans credentials.yaml")


def tickets_jql_request(username):
    return f'assignee = "{username}" AND status IN ("Implemented", "Integrated", "Closed", "Verified" ) AND updated >= startOfWeek(-{WEEKS_RANGE}) ORDER BY updated DESC'


# --- Récupération des tickets ---
def get_jira_issues(cookie, fields=DEFAULT_FIELDS):
    credentials = load_credentials()
    url = f"{credentials['jira_link']}/rest/api/2/search"
    params = {
        "jql": tickets_jql_request(credentials['jira_user']),
        "maxResults": 200,
        "fields": fields,
        "expand": "changelog",
    }

    headers = {
        "Cookie": f"JSESSIONID={cookie}",
        "Accept": "application/json"
    }

    print("🔍 Connexion à Jira (via cookie de session)...")
    response = requests.get(url, params=params, headers=headers, verify=True)

    if response.status_code != 200:
        print(f"❌ Erreur {response.status_code}: {response.text[:500]}")
        return []

    data = response.json()
    # with open("debug_jira_response.json", "w", encoding="utf-8") as f:
    #     json.dump(data, f, indent=2, ensure_ascii=False)
    issues = data.get("issues", [])
    issues.sort(key=lambda x: x["fields"]["updated"], reverse=True)
    print(f"✅ {len(issues)} tickets récupérés depuis Jira.")
    return issues


# --- Récupération de la date de changement de statut ---
def get_status_change_date(issue, target_statuses=["Implemented", "Integrated"]):
    histories = issue.get("changelog", {}).get("histories", [])

    for history in histories:
        for item in history["items"]:
            if item["field"] == "status":
                if item["toString"] in target_statuses:
                    change_date = datetime.strptime(history["created"], "%Y-%m-%dT%H:%M:%S.%f%z")
                    formatted_date = change_date.strftime("%m/%d/%Y")
                    return formatted_date

    return None


# --- Mapping sprint ---
def get_sprint_name(sprint, project="CCS2"):
    with open ("config/sdv_ccs2_pi_map.yaml", "r", encoding="utf-8") as f:
        pi_data = yaml.load(f, Loader=yaml.CSafeLoader)
    for pi in pi_data.get("PI", []):
        for sprint_data in pi.get("sprints", []):
            if project in sprint_data.keys():
                if sprint in sprint_data.values():
                    return sprint_data[project]
    return ""


# --- Component name ---
def get_component_name(ticket_title):
    if re.search(r"Diag|Conf|Concal|DRF|API", ticket_title, re.IGNORECASE):
        return "Diag/Conf"
    elif re.search(r"VHAL", ticket_title, re.IGNORECASE):
        return "VHAL"
    return "X"


# --- Formatage des tickets ---
def format_issues_to_headers(issues, headers = TICKET_HEADERS):
    formatted_issues = []
    for issue in issues:
        formatted_issue = {header: "" for header in headers}
        fields = issue["fields"]
        sprint_field = fields.get("customfield_18002") or [""]
        issue_sprint = sprint_field[0].get("value", "") if sprint_field[0] else ""

        formatted_issue["Sprint"] = get_sprint_name(issue_sprint, project="CCS2")
        formatted_issue["Date"] = get_status_change_date(issue)
        formatted_issue["DayLog"] = 7.0
        formatted_issue["Ticket"] = issue["key"]
        formatted_issue["Hours spent"] = 4.0
        formatted_issue["Title"] = fields["summary"]
        formatted_issue["Project"] = "SDV" if "SDV" in formatted_issue["Ticket"] else "CCS2"
        formatted_issue["Component"] = get_component_name(fields["summary"])
        formatted_issue["Type"] = "Bug" if fields["issuetype"]["name"] == "Bug" else "Implem"
        formatted_issue["Planned/Unplanned"] = "Unplanned" if fields["issuetype"]["name"] == "Bug" else "Planned"
        formatted_issue["Cause"] = ""
        formatted_issues.append(formatted_issue)

    return formatted_issues


# --- Export Excel ---
def export_to_excel(issues_dict: dict):
    workbook = xlsxwriter.Workbook(OUTPUT_EXCEL)
    worksheet = workbook.add_worksheet("Tickets Jira")
    format = workbook.add_format({'text_wrap': True, 'valign': 'top'})

    headers = TICKET_HEADERS
    for col, header in enumerate(headers):
        worksheet.write(0, col, header)

    for row, issue in enumerate(issues_dict, start=1):
        for col, header in enumerate(headers):
            worksheet.write(row, col, issue[header], format)

    workbook.close()
    print(f"📁 Fichier généré : {OUTPUT_EXCEL}")


# --- Trouver première ligne vide ---
def find_next_empty_row(sheet, starting_row=2, ticket_column=4):
    """
    Cherche la première ligne vide dans la colonne Ticket
    (colonne D = index 4 en Excel humain)
    """
    row = starting_row  # On suppose que la ligne 1 est le header
    while sheet.cell(row=row, column=ticket_column).value:
        row += 1
    return row


# --- Export sur Template Excel ---
def export_to_excel_template(issues_dict: dict):

    output_excel = f"{Path(EXCEL_TEMPLATE).stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    workbook = load_workbook(EXCEL_TEMPLATE)
    # format = workbook.add_format({'text_wrap': True, 'valign': 'top'})

    headers = TICKET_HEADERS
    # for col, header in enumerate(headers):
    #     worksheet.write(0, col, header)

    current_row = 2

    for row, issue in enumerate(issues_dict, start=1):
        empty_row = find_next_empty_row(workbook["Tasks"], starting_row=current_row)
        for col, header in enumerate(headers):
            workbook["Tasks"].cell(row=empty_row, column=col+1, value=issue[header])
        current_row = empty_row

    workbook.save(output_excel)
    print(f"✅ Tickets écrits dans le template Excel {output_excel}")
